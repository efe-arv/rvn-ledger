"""Synthetic-only regressions for the explain and XLSX boundary."""
import json
import tempfile
import unittest
from pathlib import Path
from rvn_ledger.excel import export_inputs, import_inputs, export_report
from rvn_ledger.diagnostics import explain, load_verified
from rvn_ledger.inputs import InputError
from rvn_ledger.outputs import OutputError


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def raw(self):
        return {'events.jsonl': b'{"event_id":"=danger", "units":9007199254740993}\r\n\xff\n\nlast',
                'accounts.json': b'[]\n', 'plans.json': b'{}\n', 'period.json': b'{}\n'}

    def test_lossless_transport_without_validation(self):
        path = self.root / 'inputs.xlsx'
        raw = self.raw()
        export_inputs(raw, path)
        self.assertEqual(import_inputs(path), raw)

    def test_export_does_not_overwrite(self):
        path = self.root / 'inputs.xlsx'
        path.write_bytes(b'original')
        with self.assertRaises(OutputError):
            export_inputs(self.raw(), path)
        self.assertEqual(path.read_bytes(), b'original')

    def test_formula_rejected_even_with_data_only_cache(self):
        from openpyxl import load_workbook
        path = self.root / 'inputs.xlsx'
        export_inputs(self.raw(), path)
        wb = load_workbook(path)
        wb['Events']['C2'] = '=1+1'
        wb.save(path)
        wb.close()
        with self.assertRaises(InputError):
            import_inputs(path)

    def test_wrong_header_rejected(self):
        from openpyxl import load_workbook
        path = self.root / 'inputs.xlsx'
        export_inputs(self.raw(), path)
        wb = load_workbook(path)
        wb['Events']['A1'] = 'wrong'
        wb.save(path)
        wb.close()
        with self.assertRaises(InputError):
            import_inputs(path)

    def test_unknown_sheet_rejected(self):
        from openpyxl import load_workbook
        path = self.root / 'inputs.xlsx'
        export_inputs(self.raw(), path)
        wb = load_workbook(path)
        wb.create_sheet('Unexpected')
        wb.save(path)
        wb.close()
        with self.assertRaises(InputError):
            import_inputs(path)

    def test_non_zip_rejected(self):
        path = self.root / 'bad.xlsx'
        path.write_bytes(b'not a workbook')
        with self.assertRaises(InputError):
            import_inputs(path)

    def test_numeric_text_rejected_no_excel_precision_coercion(self):
        from openpyxl import load_workbook
        path = self.root / 'inputs.xlsx'
        export_inputs(self.raw(), path)
        wb = load_workbook(path)
        wb['Events']['C2'] = 123
        wb.save(path)
        wb.close()
        with self.assertRaises(InputError):
            import_inputs(path)

    def test_empty_events_preserved(self):
        raw = self.raw()
        raw['events.jsonl'] = b''
        path = self.root / 'empty.xlsx'
        export_inputs(raw, path)
        self.assertEqual(import_inputs(path), raw)

    def test_report_preserves_big_numbers_and_formula_like_ids(self):
        from openpyxl import load_workbook
        path = self.root / 'report.xlsx'
        data = {'manifest': {'counts': {'invoices': 1}, 'totals_by_currency': {'USD': {'total_minor': 9007199254740993}}, 'inputs': {}},
                'invoices': [{'account_id': '=danger', 'currency': 'USD', 'timezone': 'UTC', 'subtotal_minor': 9007199254740993,
                              'credit_applied_minor': 0, 'credit_remaining_minor': 0, 'total_minor': 9007199254740993,
                              'quarantined_count': 0, 'lines': []}],
                'audit': {'decisions': [], 'invoices': {}}}
        export_report(data, path)
        wb = load_workbook(path)
        self.addCleanup(wb.close)
        self.assertEqual(wb['Invoices']['A2'].value, '=danger')
        self.assertEqual(wb['Invoices']['A2'].data_type, 's')
        self.assertEqual(wb['Invoices']['D2'].value, '9007199254740993')
        self.assertFalse(any(c.data_type == 'f' for ws in wb for row in ws for c in row))

    def test_explain_filters_and_reports_sources(self):
        data = {'manifest': {'inputs': {'events.jsonl': {'sha256': 'a' * 64}}},
                'audit': {'decisions': [{'line': 2, 'event_id': 'ev', 'status': 'quarantined', 'sha256': 'b'*64,
                                         'reasons': ['invalid_units']}]}}
        result = explain(data, event_id='ev')
        self.assertEqual(result['records'][0]['source']['line'], 2)
        self.assertEqual(result['records'][0]['reasons'][0]['code'], 'invalid_units')
        with self.assertRaises(InputError):
            explain(data, event_id='missing')
        with self.assertRaises(InputError):
            explain(data, line=0)

    def test_cli_roundtrip_explain_and_export(self):
        import contextlib
        import io
        from test_pipeline import write_fixture
        from rvn_ledger.cli import main
        inputs = write_fixture(self.root / 'original')
        out = self.root / 'out'
        workbook = self.root / 'inputs.xlsx'
        imported = self.root / 'imported'
        with contextlib.redirect_stdout(io.StringIO()) as stdout:
            self.assertEqual(main(['run', '--input-dir', str(inputs), '--out', str(out), '--json']), 0)
            stdout.seek(0); stdout.truncate()
            self.assertEqual(main(['explain', '--out', str(out), '--events', str(inputs / 'events.jsonl'), '--json']), 0)
            self.assertTrue(json.loads(stdout.getvalue())['source_values_verified'])
            self.assertEqual(main(['export', '--out', str(out), '--file', str(self.root / 'report.xlsx')]), 0)
            self.assertEqual(main(['export', '--kind', 'inputs', '--input-dir', str(inputs), '--file', str(workbook)]), 0)
            self.assertEqual(main(['import', '--file', str(workbook), '--input-dir', str(imported)]), 0)
            for name in ('events.jsonl', 'accounts.json', 'plans.json', 'period.json'):
                self.assertEqual((inputs / name).read_bytes(), (imported / name).read_bytes())
            second = self.root / 'second'
            self.assertEqual(main(['run', '--input-dir', str(imported), '--out', str(second)]), 0)
            for name in ('invoices.json', 'quarantine.json', 'audit.json', 'manifest.json'):
                self.assertEqual((out / name).read_bytes(), (second / name).read_bytes())
        with contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(main(['import', '--file', str(workbook), '--input-dir', str(imported)]), 2)
            self.assertEqual(main(['import', '--file', str(self.root / 'report.xlsx'), '--input-dir', str(self.root / 'bad')]), 2)
            self.assertFalse((self.root / 'bad').exists())
            (inputs / 'events.jsonl').write_bytes(b'changed')
            self.assertEqual(main(['explain', '--out', str(out), '--events', str(inputs / 'events.jsonl')]), 2)
            (out / 'audit.json').write_bytes(b'{}')
            self.assertEqual(main(['explain', '--out', str(out)]), 2)
            self.assertEqual(main(['export', '--out', str(out), '--file', str(self.root / 'bad.xlsx')]), 2)
            self.assertFalse((self.root / 'bad.xlsx').exists())

    def test_bad_config_not_published(self):
        import contextlib
        import io
        from rvn_ledger.cli import main
        workbook = self.root / 'bad-config.xlsx'
        export_inputs(self.raw(), workbook)
        with contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(main(['import', '--file', str(workbook), '--input-dir', str(self.root / 'target')]), 2)
        self.assertFalse((self.root / 'target').exists())

    def test_archive_expansion_limit(self):
        import zipfile
        from unittest.mock import patch
        path = self.root / 'inputs.xlsx'
        export_inputs(self.raw(), path)
        with patch('rvn_ledger.excel.MAX_EXPANDED', 1), self.assertRaises(InputError):
            import_inputs(path)

    def test_invalid_parts_and_event_line_numbers(self):
        from openpyxl import load_workbook
        for sheet, cell, value in [('Config', 'B2', 2), ('Config', 'A2', '../plans.json'), ('Events', 'A2', 2),
                                   ('Events', 'D2', 'NONE'), ('Events', 'B2', 'unknown')]:
            path = self.root / (sheet + cell + '.xlsx')
            export_inputs(self.raw(), path)
            wb = load_workbook(path)
            wb[sheet][cell] = value
            wb.save(path); wb.close()
            with self.subTest(sheet=sheet, cell=cell), self.assertRaises(InputError):
                import_inputs(path)

    def test_malformed_xml_entities_external_links_and_duplicate_members(self):
        import io
        import zipfile
        original = self.root / 'original.xlsx'
        export_inputs(self.raw(), original)
        with zipfile.ZipFile(original) as z:
            members = {n: z.read(n) for n in z.namelist()}
        variants = [
            {'xl/worksheets/sheet2.xml': b'<broken'},
            {'xl/worksheets/sheet2.xml': b'<!DOCTYPE x [<!ENTITY a "bad">]><x>&a;</x>'},
            {'xl/worksheets/_rels/sheet2.xml.rels': b"<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'><Relationship Id='r1' Type='x' Target='https://example.invalid' TargetMode='External'/></Relationships>"},
            {'xl/vbaProject.bin': b'fake-macro'},
            {'../escape': b'no'},
        ]
        for i, changes in enumerate(variants):
            path = self.root / f'bad-{i}.xlsx'
            with zipfile.ZipFile(path, 'w') as z:
                for name, value in {**members, **changes}.items():
                    z.writestr(name, value)
            with self.subTest(i=i), self.assertRaises(InputError):
                import_inputs(path)

    def test_import_caught_io_failure_cleans_own_target(self):
        import contextlib
        import io
        from unittest.mock import patch
        from test_pipeline import write_fixture
        from rvn_ledger.cli import main
        inputs = write_fixture(self.root / 'inputs')
        path = self.root / 'inputs.xlsx'
        export_inputs({n: (inputs / n).read_bytes() for n in ('events.jsonl', 'accounts.json', 'plans.json', 'period.json')}, path)
        target = self.root / 'target'
        with patch('os.replace', side_effect=OSError('test replacement failure')), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(main(['import', '--file', str(path), '--input-dir', str(target)]), 2)
        self.assertFalse(target.exists())
        self.assertFalse(list(self.root.glob('.ledger-import-*')))

    def test_missing_outputs_fail_closed(self):
        with self.assertRaises(OutputError):
            load_verified(self.root)


if __name__ == '__main__':
    unittest.main()
